from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def format_date(date_string):

    try:

        return datetime.strptime(date_string, "%Y-%m-%d").strftime("%d %B %Y")

    except Exception:

        return date_string


def add_heading(doc, title):

    heading = doc.add_heading(level=1)

    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = heading.add_run(title)

    run.bold = True

    run.font.size = Pt(18)


def add_field(doc, title, value):

    paragraph = doc.add_paragraph()

    heading = paragraph.add_run(f"{title}: ")

    heading.bold = True

    paragraph.add_run(str(value))


def create_doc(report, filename, report_title="Sprint Report"):

    doc = Document()

    section = doc.sections[0]

    section.left_margin = Pt(50)

    section.right_margin = Pt(50)

    section.top_margin = Pt(40)

    section.bottom_margin = Pt(40)

    title = doc.add_heading(level=0)

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run("SprintOps AI")

    run.bold = True

    run.font.size = Pt(22)

    subtitle = doc.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle.add_run("AI Powered Sprint Reporting").italic = True

    generated = doc.add_paragraph()

    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER

    generated.add_run("Generated On : ").bold = True

    generated.add_run(datetime.now().strftime("%d %B %Y %I:%M %p"))

    doc.add_page_break()

    add_heading(doc, report_title.replace("_", " ").title())

    for index, item in enumerate(report, start=1):

        doc.add_heading(f"Entry {index}", level=2)

        status = (
            item.get("status")
            or item.get("session_status")
            or item.get("current_status")
            or "N/A"
        )

        add_field(doc, "Date", format_date(item.get("date", "N/A")))

        add_field(doc, "Time", item.get("time", "Unknown"))

        add_field(doc, "Status", status)

        add_field(doc, "Today's Work", item.get("today_work", "N/A"))

        add_field(doc, "Blockers", item.get("blockers", "None"))

        add_field(doc, "Support Required", item.get("support", "No"))

        doc.add_paragraph("────────────────────────────────────────")

    doc.add_heading("Summary", level=1)

    total = len(report)

    blocked = 0

    on_track = 0

    for item in report:

        status = item.get("status", "").lower()

        if status == "blocked":

            blocked += 1

        elif status == "on_track":

            on_track += 1

    add_field(doc, "Total Entries", total)

    add_field(doc, "On Track", on_track)

    add_field(doc, "Blocked", blocked)

    doc.add_page_break()

    footer = doc.add_paragraph()

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer.add_run("Generated automatically by SprintOps AI").italic = True

    footer.add_run("\n")

    footer.add_run("Confidential - Internal Use Only")

    doc.save(filename)
